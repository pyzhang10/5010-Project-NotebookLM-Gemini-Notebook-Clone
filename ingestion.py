from __future__ import annotations

import ipaddress
import re
import shutil
import socket
import uuid
from pathlib import Path
from urllib.parse import urlparse

import fitz
import requests
from bs4 import BeautifulSoup
from pptx import Presentation

from config import CHUNK_OVERLAP, CHUNK_SIZE, MAX_UPLOAD_MB
from models import Chunk, utc_now
from storage import LocalStorage


SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".txt"}


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _windows(words: list[str], size: int, overlap: int) -> list[str]:
    if not words:
        return []
    result = []
    step = max(1, size - overlap)
    for start in range(0, len(words), step):
        value = " ".join(words[start : start + size]).strip()
        if value:
            result.append(value)
        if start + size >= len(words):
            break
    return result


def chunks_from_sections(
    sections: list[tuple[str, str]], source_id: str, source_name: str,
    source_type: str, url: str = ""
) -> list[Chunk]:
    chunks: list[Chunk] = []
    sequence = 0
    for location, text in sections:
        cleaned = _clean(text)
        for part in _windows(cleaned.split(), CHUNK_SIZE, CHUNK_OVERLAP):
            sequence += 1
            chunks.append(
                Chunk(
                    id=f"{source_id}-c{sequence}", source_id=source_id,
                    source_name=source_name, source_type=source_type,
                    location=location, text=part, url=url,
                )
            )
    return chunks


def extract_file(path: Path) -> tuple[list[tuple[str, str]], str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with fitz.open(path) as document:
            sections = [
                (f"Page {index + 1}", page.get_text("text"))
                for index, page in enumerate(document)
            ]
        return sections, "pdf"
    if suffix == ".pptx":
        presentation = Presentation(path)
        sections = []
        for index, slide in enumerate(presentation.slides):
            text = "\n".join(
                shape.text for shape in slide.shapes if hasattr(shape, "text")
            )
            sections.append((f"Slide {index + 1}", text))
        return sections, "pptx"
    if suffix == ".txt":
        return [("Text file", path.read_text(encoding="utf-8", errors="replace"))], "txt"
    raise ValueError("Unsupported file. Upload PDF, PPTX, or TXT.")


def _validate_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("Enter a valid HTTP or HTTPS URL.")
    addresses = socket.getaddrinfo(parsed.hostname, parsed.port or 443)
    for item in addresses:
        address = ipaddress.ip_address(item[4][0])
        if address.is_private or address.is_loopback or address.is_link_local:
            raise ValueError("Private or local network URLs are not allowed.")


def extract_url(url: str) -> tuple[list[tuple[str, str]], str]:
    _validate_public_url(url)
    response = requests.get(
        url,
        timeout=15,
        allow_redirects=True,
        headers={"User-Agent": "NotebookRAG/1.0 educational project"},
    )
    response.raise_for_status()
    # Re-check the final destination because a public URL could redirect to a
    # private address after the first validation.
    _validate_public_url(response.url)
    if "text/html" not in response.headers.get("content-type", ""):
        raise ValueError("URL must point to an HTML page.")
    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "noscript"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else url
    text = _clean(soup.get_text("\n"))
    if len(text) < 80:
        raise ValueError("No meaningful text was found on this page.")
    return [("Web page", text)], title


class IngestionService:
    def __init__(self, storage: LocalStorage, index):
        self.storage = storage
        self.index = index

    def ingest_file(self, notebook_id: str, uploaded_path: str) -> dict:
        source = Path(uploaded_path)
        if source.suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise ValueError("Unsupported file. Upload PDF, PPTX, or TXT.")
        if source.stat().st_size > MAX_UPLOAD_MB * 1024 * 1024:
            raise ValueError(f"File exceeds the {MAX_UPLOAD_MB} MB limit.")
        source_id = str(uuid.uuid4())
        destination = self.storage.path(notebook_id) / "raw" / f"{source_id}{source.suffix.lower()}"
        shutil.copy2(source, destination)
        try:
            sections, source_type = extract_file(destination)
            chunks = chunks_from_sections(
                sections, source_id, source.name, source_type
            )
            if not chunks:
                raise ValueError("No text could be extracted. Scanned PDFs require OCR.")
            extracted_path = self.storage.path(notebook_id) / "extracted" / f"{source_id}.txt"
            extracted_path.write_text(
                "\n\n".join(f"## {loc}\n{text}" for loc, text in sections),
                encoding="utf-8",
            )
            self.index.add(notebook_id, chunks)
            record = {
                "id": source_id, "name": source.name, "type": source_type,
                "chunks": len(chunks), "created_at": utc_now(), "url": "",
            }
            self.storage.add_source(notebook_id, record)
            return record
        except Exception:
            destination.unlink(missing_ok=True)
            raise

    def ingest_url(self, notebook_id: str, url: str) -> dict:
        sections, title = extract_url(url.strip())
        source_id = str(uuid.uuid4())
        chunks = chunks_from_sections(sections, source_id, title, "url", url.strip())
        self.index.add(notebook_id, chunks)
        extracted_path = self.storage.path(notebook_id) / "extracted" / f"{source_id}.txt"
        extracted_path.write_text(sections[0][1], encoding="utf-8")
        record = {
            "id": source_id, "name": title, "type": "url",
            "chunks": len(chunks), "created_at": utc_now(), "url": url.strip(),
        }
        self.storage.add_source(notebook_id, record)
        return record
